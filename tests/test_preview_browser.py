"""Opt-in regressions against the real offline app, Chromium, and KaTeX.

Run with MATHBANK_TEST_BROWSER=1 python -m pytest --noconftest
tests/test_preview_browser.py. Requires the optional agent-browser CLI and its
browser; neither is an application dependency. A source-only temporary copy
gets its own database, uploads, token, server, and browser session.
"""

import json
import os
from pathlib import Path
import shutil
import socket
import subprocess
import sys
import time
import urllib.request
import uuid

import pytest


ROOT = Path(__file__).resolve().parents[1]
pytestmark = pytest.mark.skipif(
    os.environ.get("MATHBANK_TEST_BROWSER") != "1",
    reason="opt-in actual browser regression",
)


class Browser:
    def __init__(self, executable, session):
        self.executable = executable
        self.session = session

    def command(self, *args, source=None):
        result = subprocess.run(
            [self.executable, "--session", self.session, "--json", *args],
            input=source, text=True, capture_output=True, timeout=45,
        )
        assert result.returncode == 0, result.stdout + result.stderr
        payload = json.loads(result.stdout)
        assert payload.get("success"), payload
        return payload.get("data")

    def evaluate(self, source):
        # Retain async results until CDP finishes awaiting them, including drags
        # that replace their source DOM nodes before their promise resolves.
        retained_source = "window.__mathbankTestEvaluation = (0, eval)(" + json.dumps(source) + ")"
        data = self.command("eval", "--stdin", source=retained_source)
        return data.get("result") if isinstance(data, dict) and "result" in data else data

    def settle(self):
        if self.evaluate("typeof window.MathBankBrowserChecks === 'undefined'"):
            self.evaluate((ROOT / "tests" / "browser_preview_fixture.js").read_text())
        self.evaluate("MathBankBrowserChecks.settle().then(() => true)")


def complete_import_fixture_categories(browser, index):
    """Give minimal mock questions valid required fields, without bypassing validation."""
    assert browser.evaluate("""(() => {
        const card=document.getElementById('parsed-card-' + """ + str(index) + """);
        const comp=card.querySelector('.card-compulsory');
        const chapter=card.querySelector('.card-chapter');
        const first=[...comp.options].find(option=>option.value);
        if(!first)return false;
        comp.value=first.value;comp.dispatchEvent(new Event('change',{bubbles:true}));
        const next=[...chapter.options].find(option=>option.value);
        if(!next)return false;
        chapter.value=next.value;chapter.dispatchEvent(new Event('change',{bubbles:true}));
        return Boolean(comp.value&&chapter.value);
    })()""")


@pytest.fixture(scope="module")
def browser(tmp_path_factory):
    executable = shutil.which("agent-browser")
    assert executable, "Install the optional agent-browser CLI to run this test"
    root = tmp_path_factory.mktemp("mathbank-browser")
    shutil.copy2(ROOT / "main.py", root / "main.py")
    for directory in ("mathbank", "templates", "static"):
        shutil.copytree(
            ROOT / directory, root / directory,
            ignore=shutil.ignore_patterns("__pycache__", "uploads", "test_uploads", "*.pyc"),
        )
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        port = sock.getsockname()[1]
    url = f"http://127.0.0.1:{port}"
    env = dict(os.environ, MATHBANK_LAUNCH_ID=uuid.uuid4().hex)
    env.pop("PYTHONPATH", None)
    browser = Browser(executable, "mathbank-test-" + uuid.uuid4().hex[:12])
    browser.root = root  # Disposable app root for opt-in local fixture assets.
    opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
    with (root / "server.log").open("w") as log:
        server = subprocess.Popen(
            [sys.executable, "-u", "-m", "uvicorn", "main:app", "--host", "127.0.0.1", "--port", str(port)],
            cwd=root, env=env, stdout=log, stderr=subprocess.STDOUT,
        )
        try:
            deadline = time.monotonic() + 40
            while time.monotonic() < deadline:
                try:
                    with opener.open(url + "/healthz", timeout=1) as response:
                        if json.load(response).get("ready"):
                            break
                except (OSError, ValueError):
                    pass
                if server.poll() is not None:
                    pytest.fail((root / "server.log").read_text())
                time.sleep(0.1)
            else:
                pytest.fail((root / "server.log").read_text())
            browser.command("open", url)
            browser.command("set", "viewport", "1600", "1100")
            browser.command("wait", "--fn", "typeof window.renderPaperCanvas === 'function'")
            browser.evaluate("window.selectWorkspace('paper', '组卷排版工作台'); true")
            browser.command("wait", "--fn", "!!document.getElementById('a4PaperPreviewSheet') && !window.PaperStore.cartQuestionLoad.loading")
            browser.evaluate((ROOT / "tests" / "browser_preview_fixture.js").read_text())
            yield browser
        finally:
            try:
                subprocess.run(
                    [executable, "--session", browser.session, "close"],
                    capture_output=True, timeout=20, check=False,
                )
            finally:
                server.terminate()
                try:
                    server.wait(timeout=10)
                except subprocess.TimeoutExpired:
                    server.kill()
                    server.wait(timeout=5)


def test_association_switching_and_real_unsaved_changes(browser, tmp_path):
    ids = browser.evaluate(r"""
    (async () => {
        selectWorkspace('bank', '题库研讨工作台');
        const ids = [];
        for (const content of ['关联回归甲：求 $1+2$。', '关联回归乙：求 $3+4$。']) {
            const form = new FormData();
            form.set('content', content);
            form.set('question_type', 'detailed_answer');
            form.set('difficulty', 'medium');
            const res = await fetch('/api/questions', {method: 'POST', body: form});
            const data = await res.json();
            if (!res.ok || !data.question) throw new Error(JSON.stringify(data));
            ids.push(data.question.id);
        }
        loadQuestions();
        return ids;
    })()
    """)
    first, second = ids

    def click_association_action(function_name):
        selector = f'button[onclick="{function_name}()"]'
        # The editor scrolls smoothly; settle before the coordinate-based click.
        browser.command("scrollintoview", selector)
        browser.settle()
        browser.command("click", selector)

    def open_question(question_id, related_id):
        if browser.evaluate("!document.getElementById('editorSection').classList.contains('hidden')"):
            browser.command("click", "#editorSection .question-editor-close")
            browser.command("wait", "--fn", "document.getElementById('editorSection').classList.contains('hidden')")
        browser.command("wait", "--fn", f"!!document.querySelector('#questionsList [data-id=\"{question_id}\"]')")
        browser.command("click", f'#questionsList [data-id="{question_id}"]')
        browser.command("wait", "--fn", (
            f"EditorState.questionId === {question_id} && !questionDetailLoading"
            f" && document.getElementById('editRelatedQuestion').value === '{related_id}'"
            " && !isEditorModified()"
        ))
        assert not browser.evaluate("!!document.getElementById('unsavedChangesModalTitle')")
        browser.command("click", "#editQuestionFromPreviewBtn")
        browser.command("click", '[data-editor-panel-target="answer"]')

    open_question(first, '')
    browser.command("wait", "--fn", f"!!document.querySelector('#editRelatedQuestion option[value=\"{second}\"]')")
    browser.command("select", "#editRelatedQuestion", str(second))
    assert browser.evaluate("document.getElementById('editRelatedQuestion').value") == str(second)
    click_association_action("associateRelatedQuestion")
    browser.settle()
    assert browser.evaluate(
        f"fetch('/api/questions/{first}/associated').then(r => r.json())"
        f".then(qs => qs.some(q => q.id === {second}))"
    ), browser.evaluate("({selection: document.getElementById('editRelatedQuestion').value, questionId: EditorState.questionId})")
    browser.command("wait", "--fn", "!isEditorModified() && !document.getElementById('paperAssociatedWrapper').classList.contains('hidden')")
    for question_id, related_id in [(second, first), (first, second), (second, first), (first, second)]:
        open_question(question_id, related_id)

    # Real DELETE persists and leaves the editor clean.
    click_association_action("clearRelatedQuestion")
    browser.command("wait", "--fn", "document.getElementById('editRelatedQuestion').value === '' && !isEditorModified()")
    assert browser.evaluate(f"fetch('/api/questions/{first}/associated').then(r => r.json())") == []

    # Linking must not mark a separately edited stem as saved.
    browser.command("click", '[data-editor-panel-target="content"]')
    browser.command("fill", "#editContent", "真正尚未保存的题干修改")
    browser.command("click", '[data-editor-panel-target="answer"]')
    browser.command("select", "#editRelatedQuestion", str(second))
    click_association_action("associateRelatedQuestion")
    browser.command("wait", "--fn", "!document.getElementById('paperAssociatedWrapper').classList.contains('hidden')")
    assert browser.evaluate("isEditorModified()")
    browser.command("click", "#editorSection .question-editor-close")
    browser.command("wait", "#unsavedChangesModalTitle")
    assert browser.evaluate("EditorState.questionId") == first
    browser.command("screenshot", str(tmp_path / "association-real-unsaved-prompt.png"))
    # Close the modal without saving test edits, keeping other browser tests independent.
    browser.command("click", "#cancelBtn")
    browser.command("click", '[data-editor-panel-target="content"]')
    browser.command("fill", "#editContent", "关联回归甲：求 $1+2$。")
    browser.command("click", "#editorSection .question-editor-close")
    browser.evaluate("selectWorkspace('paper', '组卷排版工作台'); true")


def test_valid_and_naked_table_formulas_render_without_internal_placeholders(browser):
    result = browser.evaluate(r"""
    (() => {
        const samples = [
            String.raw`\begin{tabular}{cc}$x_1$ & $y_1$ \\ $x_2$ & $y_2$\end{tabular}`,
            String.raw`\begin{tabular}{cc}x_1 & $y_1$ \\ $\frac{1}{2}$ & y_2\end{tabular}`,
            String.raw`\begin{tabular}{cc}\multicolumn{2}{c}{$x_1+y_1$} \\ $x_2$ & y_2\end{tabular}`,
        ];
        return samples.map(source => {
            const host = document.createElement('div');
            window.renderQuestionPreviewContent(host, source);
            return {
                cells: host.querySelectorAll('td').length,
                math: host.querySelectorAll('.katex').length,
                errors: host.querySelectorAll('.katex-error').length,
                leaked: /PLACEHOLDER|@@|\uE000/.test(host.textContent),
                tex: [...host.querySelectorAll('annotation')].map(n => n.textContent),
            };
        });
    })()
    """)
    assert [row["cells"] for row in result] == [4, 4, 3], result
    assert [row["math"] for row in result] == [4, 4, 3], result
    assert all(not row["leaked"] and not row["errors"] for row in result), result
    assert result[0]["tex"] == ["x_1", "y_1", "x_2", "y_2"]


@pytest.mark.parametrize("viewport", [(1440, 1000), (375, 812)])
def test_editor_confirmations_are_visible_and_restore_editing(browser, tmp_path, viewport):
    browser.command("set", "viewport", *map(str, viewport))
    browser.evaluate("selectWorkspace('bank', '题库管理'); true")
    browser.command("click", ".bank-create-btn")
    browser.command("wait", "--fn", "!document.getElementById('editorSection').classList.contains('hidden')")
    assert browser.evaluate("document.getElementById('editorSection').dataset.panel") == 'content'
    browser.command("click", '[data-editor-panel-target="classification"]')
    browser.command("fill", "#editSource", "未保存确认回归")
    browser.command("click", "#editorSection .question-editor-close")
    browser.command("wait", "--fn", "!!document.getElementById('unsavedChangesModalTitle')")

    def assert_button_not_visually_covered(button_id):
        # Inert elements still paint, but hit testing skips them. Include them
        # briefly so a painted editor covering a confirmation cannot pass.
        result = browser.evaluate("""
        (() => {
            const button = document.getElementById(%s);
            const inertElements = [...document.querySelectorAll('[inert]')];
            try {
                inertElements.forEach(element => { element.inert = false; });
                const rect = button.getBoundingClientRect();
                const hit = document.elementFromPoint(rect.x + rect.width / 2, rect.y + rect.height / 2);
                return {visible: button.contains(hit), hit: hit && (hit.id || hit.className)};
            } finally {
                inertElements.forEach(element => { element.inert = true; });
            }
        })()
        """ % json.dumps(button_id))
        assert result["visible"], result

    assert_button_not_visually_covered("cancelBtn")
    browser.command("screenshot", str(tmp_path / "unsaved-confirmation.png"))
    browser.command("click", "#cancelBtn")
    browser.command("wait", "--fn", "!document.getElementById('unsavedChangesModalTitle')")
    assert browser.evaluate("!document.querySelector('main').inert && isEditorModified()")
    assert browser.evaluate("document.getElementById('editSource').value") == "未保存确认回归"

    browser.command("click", "#editorSection .question-editor-close")
    browser.command("wait", "--fn", "!!document.getElementById('unsavedChangesModalTitle')")
    browser.command("press", "Escape")
    assert browser.evaluate("!document.querySelector('main').inert && !document.getElementById('editorSection').classList.contains('hidden')")

    browser.command("click", '[data-editor-panel-target="content"]')
    browser.command("fill", "#editContent", "求 1+1 的值。")
    browser.command("click", "#saveQuestionBtn")
    browser.command("wait", "--fn", "!!document.getElementById('missingClassificationModalTitle')")
    browser.command("wait", "--fn", "getComputedStyle(document.getElementById('missingClassificationModalTitle').closest('[role=dialog]')).opacity === '1'")
    assert_button_not_visually_covered("cancelCompulsoryBtn")
    browser.command("screenshot", str(tmp_path / "missing-classification.png"))
    browser.command("click", "#cancelCompulsoryBtn")
    browser.command("wait", "--fn", "!isQuestionSaveInFlight() && !document.querySelector('main').inert")
    assert browser.evaluate("document.getElementById('editContent').value") == "求 1+1 的值。"

    browser.command("click", "#editorSection .question-editor-close")
    browser.command("wait", "--fn", "!!document.getElementById('discardBtn')")
    browser.command("click", "#discardBtn")
    browser.command("wait", "--fn", "document.getElementById('editorSection').classList.contains('hidden')")
    assert browser.evaluate("!document.querySelector('main').inert && !document.getElementById('appNavigation').inert")
    browser.evaluate("selectWorkspace('paper', '智能组卷'); true")
    browser.command("set", "viewport", "1600", "1100")


def test_resizing_repartitions_every_question_inside_footer_boundary(browser):
    browser.command("set", "viewport", "1600", "1100")
    browser.evaluate("MathBankBrowserChecks.seed(15, 2); true")
    browser.settle()
    wide = browser.evaluate("MathBankBrowserChecks.measure()")
    browser.command("set", "viewport", "1100", "900")
    browser.settle()
    narrow = browser.evaluate("MathBankBrowserChecks.measure()")
    browser.command("set", "viewport", "1600", "1100")
    browser.settle()
    restored = browser.evaluate("MathBankBrowserChecks.measure()")
    for layout in (wide, narrow, restored):
        assert [q["id"] for page in layout for q in page["questions"]] == list(range(1, 16))
        assert all(q["overflow"] <= 2 for page in layout for q in page["questions"]), layout
    # Responsive columns may scale without adding a page. Returning to the
    # original viewport must restore the same complete question distribution.
    assert [[q["id"] for q in page["questions"]] for page in restored] == [
        [q["id"] for q in page["questions"]] for page in wide
    ]


def test_oversize_page_stays_complete_after_repeated_title_updates(browser):
    browser.command("set", "viewport", "1600", "1100")
    browser.evaluate("MathBankBrowserChecks.seed(1, 120); true")
    browser.settle()
    layouts = [browser.evaluate("MathBankBrowserChecks.measure()")]
    if browser.evaluate("document.querySelector('.paper-meta-details').hidden"):
        browser.command("click", "#togglePaperMetaBtn")
    for suffix in ("甲", "乙", "丙"):
        browser.command("fill", "#paperMetaTitle", "超长题检查" + suffix)
        browser.settle()
        layouts.append(browser.evaluate("MathBankBrowserChecks.measure()"))
    for layout in layouts:
        assert [q["id"] for page in layout for q in page["questions"]] == [1], layout
        assert any(page["expanded"] for page in layout), layout
        assert all(q["overflow"] <= 2 for page in layout for q in page["questions"]), layout


def test_header_typing_preserves_focus_and_second_character(browser):
    browser.evaluate("MathBankBrowserChecks.seed(4, 0, true); true")
    browser.settle()
    browser.command("click", ".canvas-meta-title")
    browser.command("press", "End")
    browser.command("press", "A")
    browser.settle()
    assert browser.evaluate("document.activeElement.matches('.canvas-meta-title')") is True
    browser.command("press", "B")
    browser.settle()
    assert browser.evaluate("document.querySelector('.canvas-meta-title').textContent") == "分页回归验证AB"


def test_dragging_uses_destination_question_order(browser):
    browser.evaluate("MathBankBrowserChecks.seed(4, 0, true); true")
    browser.settle()
    result = browser.evaluate("MathBankBrowserChecks.drag(1, 4)")
    assert result == [2, 3, 4, 1]
    browser.settle()
    result = browser.evaluate("MathBankBrowserChecks.drag(1, 2, true)")
    assert result == [1, 2, 3, 4]


def test_cross_page_drop_and_cancel_keep_the_cart_complete(browser):
    browser.evaluate("MathBankBrowserChecks.seed(15, 2); true")
    browser.settle()
    assert len(browser.evaluate("MathBankBrowserChecks.measure()")) > 1
    assert browser.evaluate("MathBankBrowserChecks.drag(1, 15, false, true)") == list(range(1, 16))
    browser.settle()
    assert browser.evaluate("MathBankBrowserChecks.drag(1, 15, false, false, true)") == list(range(2, 16)) + [1]
    browser.settle()
    layout = browser.evaluate("MathBankBrowserChecks.measure()")
    assert [q["id"] for page in layout for q in page["questions"]] == list(range(2, 16)) + [1]
    assert all(q["overflow"] <= 2 for page in layout for q in page["questions"]), layout


def test_returning_to_dragged_question_does_not_commit_old_hover_target(browser):
    browser.evaluate("MathBankBrowserChecks.seed(3, 0, true); true")
    browser.settle()
    assert browser.evaluate("MathBankBrowserChecks.drag(1, 3, false, false, false, true)") == [1, 2, 3]


@pytest.mark.parametrize("paper_type", ["exam", "quiz", "exam_19"])
def test_templates_preserve_mixed_question_types_and_solution_space(browser, paper_type):
    result = browser.evaluate(r"""
    (async () => {
        MathBankBrowserChecks.seed(19, 2);
        const state = window.PaperStore;
        state.meta.paper_type = %s;
        for (let id = 1; id <= 19; id++) {
            const q = state.questionsMap[id];
            if (id >= 9 && id <= 11) q.question_type = 'multi_choice';
            if (id >= 12 && id <= 14) {
                q.question_type = 'fill_in_blank';
                q.content = String.raw`已知 $x=1$，则 $x+1=$\fillin。`;
            }
            if (id >= 15) {
                q.question_type = 'detailed_answer';
                q.content = String.raw`已知函数 $f(x)=x^2$，求 $f(2)$ 并说明理由。`;
                state.cart[id - 1].solution_space = 3;
            }
        }
        window.renderPaperCanvas();
        await MathBankBrowserChecks.settle();
        return MathBankBrowserChecks.measure();
    })()
    """ % json.dumps(paper_type))
    assert [q["id"] for page in result for q in page["questions"]] == list(range(1, 20))
    assert all(q["overflow"] <= 2 for page in result for q in page["questions"]), result


def test_old_layout_callback_does_not_disconnect_current_observer(browser):
    result = browser.evaluate("""
    (() => {
        MathBankBrowserChecks.seed(15, 2);
        const oldCallback = window.scheduleActiveA4Repagination;
        window.renderPaperCanvas();
        const currentObserver = window.activeA4PaginationResizeObserver;
        oldCallback();
        return !!currentObserver && window.activeA4PaginationResizeObserver === currentObserver;
    })()
    """)
    assert result is True


def test_record_search_and_narrow_workspace_navigation(browser, tmp_path):
    browser.evaluate(r"""
    (async () => {
        const form = new FormData();
        form.set('content', '界面回归题：求 $2+3$。');
        form.set('question_type', 'detailed_answer');
        form.set('difficulty', 'medium');
        form.set('answer_markdown', '$5$');
        const question = await fetch('/api/questions', {method:'POST', body:form}).then(r=>r.json());
        if (!question.question) throw new Error('Cannot seed question');
        for (const [title, paper_type] of [['界面回归甲', 'exam'], ['界面回归乙', 'quiz']]) {
            const result = await fetch('/api/paper/save', {method:'POST', headers:{'Content-Type':'application/json'},
                body:JSON.stringify({title, paper_type, questions:[{id:question.question.id,score:5}]})}).then(r=>r.json());
            if (result.status !== 'success') throw new Error('Cannot seed paper');
        }
        selectWorkspace('dashboard', '工作台');
        loadQuestions();
        return true;
    })()
    """)
    browser.command('click', '.dashboard-link-action')
    browser.command('wait', '--fn', "document.querySelectorAll('.saved-paper-card').length === 2")
    browser.command('fill', '#savedPaperSearch', '界面回归甲')
    assert browser.evaluate("document.querySelectorAll('.saved-paper-card').length") == 1
    browser.command('select', '#savedPaperType', 'quiz')
    assert browser.evaluate("document.getElementById('savedPapersListContainer').innerText.includes('没有匹配')")
    browser.command('fill', '#savedPaperSearch', '')
    assert browser.evaluate("document.querySelectorAll('.saved-paper-card').length") == 1
    browser.command('select', '#savedPaperType', '')
    browser.command('select', '#savedPaperSort', 'oldest')
    assert browser.evaluate("document.querySelector('.saved-paper-card h4').textContent") == '界面回归甲'
    browser.command('click', '#appNavDashboard')
    browser.command('click', '.dashboard-link-action')
    browser.command('wait', '--fn', "document.querySelectorAll('.saved-paper-card').length === 2")

    browser.command('set', 'viewport', '375', '812')
    browser.command('click', '#appNavDashboard')
    browser.command('screenshot', str(tmp_path/'mobile-home.png'))
    browser.command('click', '#appNavBank')
    assert not browser.evaluate("document.getElementById('bankFilterFields').getClientRects().length > 0")
    browser.command('click', '.bank-filter-toggle')
    assert browser.evaluate("document.getElementById('bankFilterFields').getClientRects().length > 0")
    browser.command('click', '.bank-filter-toggle')
    browser.command('wait', '--fn', "!!document.querySelector('#questionsList .question-card')")
    browser.command('click', '#questionsList .question-card:first-child')
    browser.command('wait', '--fn', "document.getElementById('bankWorkspaceSection').classList.contains('bank-detail-open')")
    assert browser.evaluate("document.getElementById('previewSection').getClientRects().length > 0")
    assert not browser.evaluate("document.getElementById('sidebarSection').getClientRects().length > 0")
    browser.command('screenshot', str(tmp_path/'mobile-detail.png'))
    browser.command('click', '.bank-detail-back')
    assert browser.evaluate("document.getElementById('sidebarSection').getClientRects().length > 0")
    browser.command('set', 'viewport', '1600', '1100')


def test_import_disclosure_follows_file_type_and_reset(browser, tmp_path):
    import pymupdf
    from docx import Document

    pdf = tmp_path / 'preview.pdf'
    with pymupdf.open() as document:
        document.new_page().insert_text((72, 72), 'Question 1: 2 + 3 = ?')
        document.save(pdf)
    word = tmp_path / 'preview.docx'
    document = Document()
    document.add_paragraph('Question 1: 2 + 3 = ?')
    document.save(word)
    tex = tmp_path / 'preview.tex'
    tex.write_text(r'\documentclass{article}\begin{document}Question 1: $2+3$\end{document}')

    browser.command('click', '#appNavImport')
    assert browser.evaluate("!document.getElementById('importSourceDetails').open && document.getElementById('texImagesSection').hidden")
    browser.command('upload', '#texFileInput', str(pdf))
    assert browser.evaluate("document.getElementById('importSourceDetails').hidden && document.getElementById('texImagesSection').hidden")
    assert browser.evaluate("!document.getElementById('pdfPageRangeContainer').classList.contains('hidden')")
    browser.command('upload', '#texFileInput', str(word))
    assert browser.evaluate("document.getElementById('importFileSummary').textContent.includes('Word') && document.getElementById('pdfPageRangeContainer').classList.contains('hidden')")
    browser.command('upload', '#texFileInput', str(tex))
    browser.command('wait', '--fn', "!document.getElementById('importLatexContent').disabled && document.getElementById('importLatexContent').value.includes('Question 1')")
    assert browser.evaluate("document.getElementById('importSourceDetails').open && !document.getElementById('texImagesSection').hidden")
    # The native confirmation is unchanged; exercise reset after acceptance.
    browser.evaluate('clearAllImportInputs(); true')
    browser.command('wait', '--fn', "!document.getElementById('importLatexContent').value")
    assert browser.evaluate("!document.getElementById('importSourceDetails').open && document.getElementById('texImagesSection').hidden")


def test_import_source_review_is_optional_and_retains_evidence(browser, tmp_path):
    browser.evaluate(r"""
    (() => {
        selectWorkspace('import', '导入中心');
        replaceParsedQuestions([
            {content: '计算 $1+1$。', answer_markdown: '', question_type: 'detailed_answer'},
            {content: '已知 $x^2-1$，求最小值。', answer_markdown: '', question_type: 'detailed_answer',
             source_review: {required: true, reasons: ['公式与原文不同，请核对正负号。'],
                source_excerpt: '2. 已知 $x^2+1$，求最小值。<img src=x onerror=window.reviewInjected=true>'}}
        ]);
        renderParsedQuestionsList(parsedQuestionsData);
        renderSourceIntegrityReport({source_review_count: 1, unmatched_source: [
            {source_excerpt: '3. 计算 $3+4$。', reason: '原文题目未对应到结果'}
        ]});
        document.getElementById('importPlaceholder').classList.add('hidden');
        document.getElementById('parsedQuestionsWrapper').classList.remove('hidden');
        return true;
    })()
    """)
    browser.command('wait', '--fn', "!!document.querySelector('#parsed-card-1 .card-source-review-panel')")
    assert browser.evaluate("document.querySelectorAll('.card-source-review-confirm').length") == 0
    assert browser.evaluate("!document.querySelector('#parsed-card-1 .card-source-review-panel').open")
    assert browser.evaluate("!document.getElementById('parsedSourceIntegrityReport').open")
    assert browser.evaluate('getCheckedUnsavedIndices()') == [0, 1]
    assert not browser.evaluate("document.querySelector('#parsed-card-1 .card-select-checkbox').disabled")
    for selector in ('#parsed-card-1 .card-source-review-panel > summary',
                     '#parsed-card-1 .card-source-review-panel > details > summary'):
        browser.command('scrollintoview', selector)
        browser.settle()
        browser.command('click', selector)
        browser.settle()
    source_state = browser.evaluate("({outerOpen:document.querySelector('#parsed-card-1 .card-source-review-panel').open,sourceOpen:document.querySelector('#parsed-card-1 .card-source-review-panel > details').open,mathRendered:document.querySelector('#parsed-card-1 .card-source-review-panel > details').dataset.mathRendered,renderer:typeof renderMathInElement,katex:!!document.querySelector('#parsed-card-1 .card-source-review-panel .katex')})")
    (tmp_path / 'source-review-toggle-state.json').write_text(json.dumps(source_state, ensure_ascii=False))
    assert source_state['outerOpen'] and source_state['sourceOpen'] and source_state['katex'], source_state
    assert browser.evaluate("typeof window.reviewInjected === 'undefined'")
    assert browser.evaluate("document.querySelectorAll('#parsed-card-1 details img').length") == 0
    assert '3. 计算' in browser.evaluate("document.getElementById('parsedSourceIntegrityReport').textContent")
    assert browser.evaluate("toggleSelectAllParsed(true); getCheckedUnsavedIndices()") == [0, 1]
    complete_import_fixture_categories(browser, 1)
    assert browser.evaluate("validateParsedQuestionBeforeImport(1,{notify:false,focus:false})")
    assert browser.evaluate('parsedQuestionNeedsSourceReview(1)')
    browser.evaluate(r"""
    (() => {
        const input = document.querySelector('#parsed-card-1 .card-content-textarea');
        input.value += ' 请说明理由。';
        input.dispatchEvent(new Event('input', {bubbles: true}));
        return true;
    })()
    """)
    assert browser.evaluate('parsedQuestionNeedsSourceReview(1)')
    assert browser.evaluate('parsedQuestionsData[1].source_review.required') is True
    assert browser.evaluate('getCheckedUnsavedIndices()') == [0, 1]
    assert browser.evaluate("document.querySelectorAll('.card-source-review-confirm').length") == 0
    browser.command('screenshot', str(tmp_path / 'import-source-review-optional.png'))
    browser.evaluate("replaceParsedQuestions([]); renderParsedQuestionsList([]); true")
    assert browser.evaluate("!document.getElementById('parsedSourceIntegrityReport')")


def test_formula_choices_keep_each_option_on_one_line(browser):
    result = browser.evaluate(r"""
    (() => {
        const previous = document.getElementById('formulaChoiceNoWrapFixture');
        if (previous) previous.remove();
        const host = document.createElement('section');
        host.id = 'formulaChoiceNoWrapFixture';
        host.style.cssText = 'position:fixed;left:0;top:0;width:520px;padding:8px;background:white;z-index:99999';
        document.body.appendChild(host);
        const source = String.raw`选择正确答案。\begin{choices}
\item $y=(x-888888)^2-600000$
\item $y=(x-888888)^2+200000$
\item $y=(x+222222)^2-600000$
\item $y=(x+222222)^2+200000$
\end{choices}`;
        renderQuestionPreviewContent(host, source);
        adaptChoicesGridLayout(host);
        const grid = host.querySelector('.choices-grid');
        const items = [...grid.querySelectorAll('.choices-content')].map(content => {
            const renderedParts = [...content.querySelectorAll('.katex .base')];
            const lineTops = [...new Set(renderedParts.map(part => Math.round(part.getBoundingClientRect().top)))];
            return {
                whiteSpace: getComputedStyle(content).whiteSpace,
                renderedParts: renderedParts.length,
                lineCount: lineTops.length,
                overflow: content.scrollWidth > content.clientWidth + 1,
            };
        });
        const response = {columns: grid.dataset.choiceColumns, items};
        host.remove();
        return response;
    })()
    """)
    assert result['columns'] in ('1', '2'), result
    assert all(item['whiteSpace'] == 'nowrap' for item in result['items']), result
    assert all(item['renderedParts'] > 0 and item['lineCount'] == 1 for item in result['items']), result
    assert not any(item['overflow'] for item in result['items']), result


def test_image_options_stay_in_labeled_import_editor_and_paper_cells(browser):
    paths = browser.evaluate(r"""
    (async () => {
        const paths = [];
        for (const label of ['A', 'B', 'C', 'D']) {
            const canvas = document.createElement('canvas');
            canvas.width = 220; canvas.height = 180;
            const context = canvas.getContext('2d');
            context.fillStyle = '#fff'; context.fillRect(0, 0, 220, 180);
            context.fillStyle = '#000'; context.font = '48px serif'; context.fillText(label, 80, 100);
            const form = new FormData();
            form.set('file', await new Promise(resolve => canvas.toBlob(resolve)), label + '.png');
            const response = await fetch('/api/upload', {method: 'POST', body: form});
            const data = await response.json();
            if (!response.ok) throw Error(JSON.stringify(data));
            paths.push(data.file_path);
        }
        selectWorkspace('import', '导入中心');
        const content = String.raw`选择正确图象 $y=x^2$。\begin{choices}` +
            paths.map(path => String.raw`\item ![](` + path + ')').join('\n') + String.raw`\end{choices}`;
        const question = {content, answer_markdown: '', image_paths: paths, question_type: 'single_choice'};
        replaceParsedQuestions([question]);
        renderParsedQuestionsList(parsedQuestionsData);
        document.getElementById('importPlaceholder').classList.add('hidden');
        document.getElementById('parsedQuestionsWrapper').classList.remove('hidden');
        window.imageOptionBrowserQuestion = question;
        return paths;
    })()
    """)
    browser.command('wait', '--fn', "[...document.querySelectorAll('#parsed-card-0 .choices-content img')].filter(i => i.complete && i.naturalWidth).length === 4")
    try:
        for width, height in [(1800, 1100), (375, 812)]:
            browser.command('set', 'viewport', str(width), str(height))
            browser.settle()
            result = browser.evaluate(r"""
            (() => {
                const host = document.querySelector('#parsed-card-0 .card-content-preview');
                adaptChoicesGridLayout(host);
                return {
                    totalImages: host.querySelectorAll('img').length,
                    columns: host.querySelector('.choices-grid').dataset.choiceColumns,
                    allVisible: [...host.querySelectorAll('img')].every(image =>
                        image.getBoundingClientRect().bottom <= host.parentElement.getBoundingClientRect().bottom + 1
                    ),
                    items: [...host.querySelectorAll('.choices-item')].map(item => {
                        const image = item.querySelector('img');
                        const content = item.querySelector('.choices-content');
                        const bounds = content.getBoundingClientRect(), imageBounds = image.getBoundingClientRect();
                        return {label: item.querySelector('.choices-label').textContent, path: image.getAttribute('src'),
                            fits: imageBounds.left >= bounds.left - 1 && imageBounds.right <= bounds.right + 1,
                            overflow: content.scrollWidth > content.clientWidth + 1};
                    })
                };
            })()
            """)
            assert result['totalImages'] == 4, result
            assert [item['label'] for item in result['items']] == ['A.', 'B.', 'C.', 'D.']
            assert [item['path'] for item in result['items']] == paths
            assert all(item['fits'] and not item['overflow'] for item in result['items']), result
            assert result['columns'] == ('2' if width == 1800 else '1'), result
            if width == 1800:
                assert result['allVisible'], result
        browser.command('set', 'viewport', '1600', '1100')
        # Save and reload through the actual API in the isolated test database,
        # then exercise the real bank detail and editor, not just the formatter.
        question_id = browser.evaluate(r"""
        (async () => {
            const question = window.imageOptionBrowserQuestion;
            const form = new FormData();
            form.set('content', question.content);
            form.set('question_type', 'single_choice');
            form.set('difficulty', 'medium');
            form.set('image_paths', JSON.stringify(question.image_paths));
            const response = await fetch('/api/questions', {method:'POST', body:form});
            const data = await response.json();
            if (!response.ok) throw Error(JSON.stringify(data));
            selectWorkspace('bank', '题库管理');
            loadQuestions();
            selectQuestion(data.question, {silent:true});
            return data.question.id;
        })()
        """)
        browser.command('wait', '--fn', f'EditorState.questionId === {question_id} && !questionDetailLoading')
        browser.command('wait', '--fn', "[...document.querySelectorAll('#paperContent .choices-item img')].filter(i=>i.complete && i.naturalWidth).length === 4")
        for selector in ('#paperContent', '#contentPreview'):
            if selector == '#contentPreview':
                browser.command('click', '#editQuestionFromPreviewBtn')
                browser.command('click', '[data-editor-panel-target="content"]')
                browser.settle()
            result = browser.evaluate(f"""
            [...document.querySelectorAll('{selector} .choices-item')].map(item => ({{
                label: item.querySelector('.choices-label').textContent,
                path: item.querySelector('img').getAttribute('src'),
                overflow: item.scrollWidth > item.clientWidth + 1
            }}))
            """)
            assert [item['label'] for item in result] == ['A.', 'B.', 'C.', 'D.']
            assert [item['path'] for item in result] == paths
            assert not any(item['overflow'] for item in result), result
        browser.command('click', '#editorSection .question-editor-close')
        browser.command('wait', '--fn', "document.getElementById('editorSection').classList.contains('hidden')")
        browser.evaluate("selectWorkspace('paper', '组卷排版工作台'); true")
        browser.command('wait', '--fn', '!PaperStore.cartQuestionLoad.loading')
        browser.evaluate(r"""
        MathBankBrowserChecks.seed(1, 0, true);
        Object.assign(PaperStore.questionsMap[1], window.imageOptionBrowserQuestion);
        renderPaperCanvas(); true;
        """)
        browser.settle()
        assert browser.evaluate("document.querySelector('#a4PaperPreviewSheet .choices-grid').dataset.choiceColumns") == '4'
        result = browser.evaluate(r"""
        [...document.querySelectorAll('#a4PaperPreviewSheet .choices-item')].map(item => ({
            label: item.querySelector('.choices-label').textContent,
            path: item.querySelector('img').getAttribute('src')
        }))
        """)
        assert [item['label'] for item in result] == ['A.', 'B.', 'C.', 'D.']
        assert [item['path'] for item in result] == paths
    finally:
        browser.command('set', 'viewport', '1600', '1100')
        browser.evaluate("replaceParsedQuestions([]); selectWorkspace('paper', '组卷排版工作台'); true")


def test_content_area_and_three_level_bank_filter(browser):
    browser.command('set', 'viewport', '1440', '1000')
    sample = browser.evaluate(r"""
    (async () => {
        const phase = Object.keys(categoryTree)[0];
        const chapters = Object.keys(categoryTree[phase]);
        const first = categoryTree[phase][chapters[0]];
        const second = categoryTree[phase][chapters[1]];
        const rows = [[chapters[0], first[0]], [chapters[0], first[1]], [chapters[1], second[0]]];
        const ids = [];
        for (let i=0; i<rows.length; i++) {
            const form = new FormData();
            Object.entries({content:'小节过滤回归 '+i, question_type:'detailed_answer', difficulty:'medium',
                category_compulsory:phase, category_chapter:rows[i][0], category_knowledge:rows[i][1]}).forEach(([k,v])=>form.set(k,v));
            const data = await fetch('/api/questions',{method:'POST',body:form}).then(r=>r.json());
            if (!data.question) throw new Error(JSON.stringify(data));
            ids.push(data.question.id);
        }
        selectWorkspace('bank','题库管理');
        return {phase, chapters, knowledge:first, secondKnowledge:second[0], ids};
    })()
    """)
    browser.command('wait', '--fn', "!document.getElementById('questionsList').hasAttribute('aria-busy')")
    assert browser.evaluate("document.getElementById('questionsList').getBoundingClientRect().height / innerHeight") > 0.65
    assert browser.evaluate("document.getElementById('questionsList').getBoundingClientRect().top / innerHeight") < 0.25
    browser.command('fill', '#searchInput', '小节过滤回归')
    browser.command('click', '.bank-filter-toggle')
    browser.command('select', '#filterCompulsory', sample['phase'])
    browser.command('select', '#filterChapter', sample['chapters'][0])
    browser.command('select', '#filterKnowledge', sample['knowledge'][0])
    browser.command('wait', '--fn', f"!document.getElementById('questionsList').hasAttribute('aria-busy') && document.querySelectorAll('#questionsList .question-card').length === 1 && document.querySelector('#questionsList .question-card').dataset.id === '{sample['ids'][0]}'")
    browser.command('select', '#filterKnowledge', sample['knowledge'][1])
    browser.command('wait', '--fn', f"!document.getElementById('questionsList').hasAttribute('aria-busy') && document.querySelector('#questionsList .question-card')?.dataset.id === '{sample['ids'][1]}'")
    saved_drafts = browser.evaluate("localStorage.getItem('mathbank_local_drafts')")
    drafts = [{"id": "section-draft-" + str(i), "isDraft": True, "content": "小节过滤回归草稿 " + str(i),
               "question_type": "detailed_answer", "difficulty": "medium", "category_compulsory": sample['phase'],
               "category_chapter": sample['chapters'][0], "category_knowledge": sample['knowledge'][i],
               "updated_at": "2026-09-18T00:00:00Z"} for i in (0, 1)]
    browser.evaluate("localStorage.setItem('mathbank_local_drafts', %s); true" % json.dumps(json.dumps(drafts)))
    browser.command('click', '#sidebarTab-drafts')
    assert browser.evaluate("document.getElementById('questionsList').innerText.includes('小节过滤回归草稿 1') && !document.getElementById('questionsList').innerText.includes('小节过滤回归草稿 0')")
    browser.command('select', '#filterKnowledge', sample['knowledge'][0])
    assert browser.evaluate("document.getElementById('questionsList').innerText.includes('小节过滤回归草稿 0') && !document.getElementById('questionsList').innerText.includes('小节过滤回归草稿 1')")
    if saved_drafts is None:
        browser.evaluate("localStorage.removeItem('mathbank_local_drafts'); true")
    else:
        browser.evaluate("localStorage.setItem('mathbank_local_drafts', %s); true" % json.dumps(saved_drafts))
    browser.command('click', '#sidebarTab-bank')
    browser.command('select', '#filterChapter', sample['chapters'][1])
    assert browser.evaluate("document.getElementById('filterKnowledge').value") == ''
    assert sample['knowledge'][0] not in browser.evaluate("Array.from(document.getElementById('filterKnowledge').options).map(o=>o.value)")
    browser.command('select', '#filterKnowledge', sample['secondKnowledge'])
    browser.command('wait', '--fn', f"!document.getElementById('questionsList').hasAttribute('aria-busy') && document.querySelector('#questionsList .question-card')?.dataset.id === '{sample['ids'][2]}'")
    browser.command('select', '#filterCompulsory', '')
    assert browser.evaluate("document.getElementById('filterChapter').value === '' && document.getElementById('filterKnowledge').value === '' && document.getElementById('filterKnowledge').disabled")
    browser.command('fill', '#searchInput', '')
    browser.command('click', '.bank-filter-toggle')
    assert browser.evaluate("document.getElementById('sidebarSection').getClientRects().length > 0")


def test_reload_returns_home_without_losing_cart_metadata_or_drafts(browser):
    saved = browser.evaluate(r"""
    (async () => {
        const questions = await fetch('/api/questions').then(r=>r.json());
        let question = questions[0];
        if (!question) {
            const form = new FormData();
            Object.entries({content:'刷新保留回归题',question_type:'detailed_answer',difficulty:'medium'}).forEach(([k,v])=>form.set(k,v));
            question = (await fetch('/api/questions',{method:'POST',body:form}).then(r=>r.json())).question;
        }
        PaperStore.cart = [];
        PaperStore.questionsMap[question.id] = question;
        addToCart(question.id, 8);
        updatePaperMeta('title','刷新保留的演示卷');
        const draft = [{id:'reload-draft',content:'刷新保留的草稿',isDraft:true}];
        localStorage.setItem('mathbank_local_drafts',JSON.stringify(draft));
        selectWorkspace('paper','智能组卷');
        return {id:question.id, draft:JSON.stringify(draft)};
    })()
    """)
    browser.command('wait', '--fn', "!!document.getElementById('togglePaperMetaBtn') && !PaperStore.cartQuestionLoad.loading")
    assert browser.evaluate("document.getElementById('paperQuestionStream').getBoundingClientRect().height / innerHeight") > 0.65
    browser.command('click', '#togglePaperFilterBtn')
    assert browser.evaluate("!document.getElementById('paperAdvancedFilters').hidden")
    browser.command('click', '#togglePaperFilterBtn')
    if browser.evaluate("!document.getElementById('paperMetaDetails').hidden"):
        browser.command('click', '#togglePaperMetaBtn')
    browser.command('click', '#togglePaperMetaBtn')
    assert browser.evaluate("!document.getElementById('paperMetaDetails').hidden")
    browser.command('click', '#togglePaperMetaBtn')
    browser.command('set', 'viewport', '375', '812')
    assert browser.evaluate("document.querySelector('.paper-config-panel').getBoundingClientRect().height / innerHeight") < 0.25
    assert browser.evaluate("document.getElementById('paperQuestionStream').getBoundingClientRect().height / innerHeight") > 0.5
    browser.command('click', '[data-paper-pane="preview"]')
    assert browser.evaluate("document.querySelector('.paper-preview-column').getClientRects().length > 0 && !document.querySelector('.paper-library-column').getClientRects().length")
    browser.command('click', '[data-paper-pane="questions"]')
    browser.command('reload')
    browser.command('wait', '--fn', "typeof PaperStore !== 'undefined' && PaperStore.activeWorkspace === 'dashboard'")
    assert browser.evaluate("!document.getElementById('dashboardWorkspaceSection').classList.contains('hidden')")
    assert browser.evaluate("PaperStore.cart.map(q=>q.id)") == [saved['id']]
    assert browser.evaluate("PaperStore.meta.title") == '刷新保留的演示卷'
    assert browser.evaluate("localStorage.getItem('mathbank_local_drafts')") == saved['draft']


def test_paper_reload_recovers_missing_selection_from_preview_button(browser):
    browser.command('set', 'viewport', '1600', '1100')
    browser.evaluate("selectWorkspace('paper', '组卷排版工作台'); true")
    browser.command('wait', '--fn', "PaperStore.activeWorkspace === 'paper' && document.getElementById('paperCanvasSection').getClientRects().length > 0 && !PaperStore.cartQuestionLoad.loading")
    fixture = browser.evaluate(r"""
    (async () => {
        const questions = [];
        for (let i = 1; i <= 6; i++) {
            const form = new FormData();
            form.set('content', `重新加载回归题 ${i}：计算 $${i}+1$。`);
            form.set('question_type', 'detailed_answer');
            form.set('difficulty', 'medium');
            const response = await fetch('/api/questions', {method:'POST',body:form});
            if (!response.ok) throw new Error('fixture question creation failed');
            questions.push((await response.json()).question);
        }
        const deletedId = questions[5].id;
        const deletion = await fetch('/api/questions/' + deletedId, {method:'DELETE'});
        if (!deletion.ok) throw new Error('fixture question deletion failed');
        PaperStore.cart = questions.map((q, i) => ({id:q.id,score:i+3,solution_space:'0'}));
        questions.forEach(q => { PaperStore.questionsMap[q.id] = q; });
        PaperStore.filters.tab = 'all';
        localStorage.setItem('mathbank_paper_cart', JSON.stringify(PaperStore.cart));
        const draft = JSON.stringify([{id:'reload-recovery-draft',content:'保留未保存草稿'}]);
        localStorage.setItem('mathbank_local_drafts', draft);
        await renderPaperWorkspace();
        window.__reloadOriginalConfirm = window.confirm;
        window.__reloadConfirmCalls = 0;
        window.confirm = () => { window.__reloadConfirmCalls++; return true; };
        return {cart:PaperStore.cart.slice(0,5),deletedId,draft};
    })()
    """)
    try:
        assert browser.evaluate("PaperStore.cartQuestionLoad.confirmedMissingIds") == [fixture['deletedId']]
        assert browser.evaluate("PaperStore.cart.length") == 6
        assert browser.evaluate("document.getElementById('paperCanvasSection').textContent.includes('卷面题目尚未完整加载')")
        assert browser.evaluate("[...document.querySelectorAll('#paperCanvasSection .paper-export-actions button')].every(button => button.disabled)")

        browser.command('click', '#paperCanvasSection button[onclick="window.retryPaperCartQuestions()"]')
        browser.command('wait', '--fn', "!PaperStore.cartQuestionLoad.loading && PaperStore.cart.length === 5")
        browser.settle()
        assert browser.evaluate("PaperStore.cart") == fixture['cart']
        assert browser.evaluate("JSON.parse(localStorage.getItem('mathbank_paper_cart'))") == fixture['cart']
        assert browser.evaluate("localStorage.getItem('mathbank_local_drafts')") == fixture['draft']
        assert browser.evaluate("window.__reloadConfirmCalls") == 0
        assert browser.evaluate("PaperStore.cartQuestionLoad.error") == ''
        assert browser.evaluate("!document.getElementById('paperCanvasSection').textContent.includes('卷面题目尚未完整加载')")
        assert browser.evaluate("[...document.querySelectorAll('#paperCanvasSection .paper-export-actions button')].every(button => !button.disabled)")
        assert browser.evaluate("document.querySelectorAll('#a4PaperPreviewSheet [data-qid]').length > 0")
    finally:
        browser.evaluate("window.confirm = window.__reloadOriginalConfirm; delete window.__reloadOriginalConfirm; true")


def test_paper_answers_preserve_reading_position_focus_and_late_page_changes(browser):
    browser.command('set', 'viewport', '1600', '1000')
    browser.evaluate((ROOT / 'tests' / 'browser_preview_fixture.js').read_text())
    browser.evaluate(r"""
    (async () => {
        selectWorkspace('paper', '组卷排版工作台');
        await MathBankBrowserChecks.settle();
        MathBankBrowserChecks.seed(24, 2, true);
        const store = PaperStore;
        store.answerCache = Object.create(null);
        store.expandedAnswerIds.clear();
        store.answerLoadingIds.clear();
        store.answerErrors = Object.create(null);
        Object.values(store.questionsMap).forEach(q => { q.has_answer = true; });
        store.questionsMap[8].content += '$$\\text{' + 'ABCD'.repeat(60) + '}$$';
        store.bankQuestions = Object.values(store.questionsMap).slice(0, 15);
        Object.assign(store.streamPagination.all, {page:1, total:24, totalPages:2, loading:false, error:''});
        await switchPaperStreamTab('all');
        const originalFetch = window.fetch;
        window.paperAnswerTest = { originalFetch, requests: [], responses: {} };
        window.fetch = function(url, ...args) {
            const match = String(url).match(/^\/api\/questions\/(8|9|10)$/);
            if (!match) return originalFetch.call(this, url, ...args);
            const id = Number(match[1]);
            paperAnswerTest.requests.push(id);
            return new Promise(resolve => { paperAnswerTest.responses[id] = resolve; });
        };
        return true;
    })()
    """)
    browser.settle()
    try:
        before = browser.evaluate(r"""
        (() => {
            const stream = document.getElementById('paperQuestionStream');
            const button = document.querySelector('button[onclick="window.togglePaperQuestionAnswer(8)"]');
            const stem = document.getElementById('paper-q-render-8');
            const card = stem.closest('.paper-resource-card');
            const score = card.querySelector('input');
            const sheet = document.getElementById('a4PaperPreviewSheet');
            stream.scrollTop += button.getBoundingClientRect().top - stream.getBoundingClientRect().top - 80;
            const reading = stem.querySelector('.katex-display') || stem;
            reading.scrollLeft = 60;
            sheet.scrollTop = 150;
            Object.assign(paperAnswerTest, {stream, button, stem, reading, card, score, sheet});
            return {scroll:stream.scrollTop, stem:reading.scrollLeft, sheet:sheet.scrollTop};
        })()
        """)
        assert before['scroll'] > 300 and before['stem'] > 0
        assert browser.evaluate("(() => { const card = paperAnswerTest.card; const title = card.querySelector('.paper-resource-heading').getBoundingClientRect(); const actions = card.querySelector('.paper-resource-actions').getBoundingClientRect(); return Math.abs((title.top + title.bottom) - (actions.top + actions.bottom)) < 4; })()")
        browser.command('click', 'button[onclick="window.togglePaperQuestionAnswer(8)"]')
        during = browser.evaluate(r"""
        (() => {
            const t = paperAnswerTest;
            return {scroll:t.stream.scrollTop, stem:document.querySelector('#paper-q-render-8 .katex-display').scrollLeft,
                sameStem:t.stem === document.getElementById('paper-q-render-8'),
                sameButton:t.button.isConnected, focused:document.activeElement === t.button,
                loading:PaperStore.answerLoadingIds.has(8), sheet:t.sheet.scrollTop};
        })()
        """)
        assert during['sameStem'] and during['sameButton'] and during['focused'], during
        assert during['loading'] and abs(during['scroll'] - before['scroll']) < 2, during
        assert during['stem'] == before['stem'] and during['sheet'] == before['sheet'], during

        # Keep typing and browsing while the answer request is still pending.
        browser.command('fill', '.paper-resource-card:has(#paper-q-render-8) input', '17')
        waiting = browser.evaluate(r"""
        (() => {
            paperAnswerTest.stream.scrollTop += 80;
            paperAnswerTest.responses[8](new Response(JSON.stringify({answer_markdown:'答案为 $x=2$。'}), {status:200}));
            return paperAnswerTest.stream.scrollTop;
        })()
        """)
        browser.command('wait', '--fn', "!PaperStore.answerLoadingIds.has(8)")
        browser.settle()
        after = browser.evaluate(r"""
        (() => {
            const t = paperAnswerTest;
            return {scroll:t.stream.scrollTop, score:t.score.value, focused:document.activeElement === t.score,
                stem:t.reading.scrollLeft, sameStem:t.stem.isConnected, sheet:t.sheet.scrollTop,
                math:!!document.querySelector('#paper-q-answer-content-8 .katex')};
        })()
        """)
        assert after['sameStem'] and after['focused'] and after['score'] == '17', after
        assert abs(after['scroll'] - waiting) < 2 and after['stem'] == before['stem'], after
        assert after['sheet'] == before['sheet'] and after['math'], after
        browser.evaluate('togglePaperQuestionAnswer(8); togglePaperQuestionAnswer(8); true')
        assert browser.evaluate('paperAnswerTest.requests.filter(id=>id===8).length') == 1
        assert browser.evaluate("paperAnswerTest.stem.isConnected && paperAnswerTest.score.value === '17'")

        # Failure and retry stay local; a collapsed pending answer stays collapsed.
        browser.evaluate('togglePaperQuestionAnswer(9); paperAnswerTest.responses[9](new Response("",{status:503})); true')
        browser.command('wait', '--fn', '!PaperStore.answerLoadingIds.has(9)')
        assert browser.evaluate("document.getElementById('paper-q-answer-9').textContent.includes('重试') && paperAnswerTest.stem.isConnected")
        browser.evaluate('retryPaperQuestionAnswer(9); togglePaperQuestionAnswer(9); true')
        browser.evaluate('togglePaperQuestionAnswer(9); true')
        assert browser.evaluate("!document.getElementById('paper-q-answer-9').hidden && PaperStore.answerLoadingIds.has(9)")
        browser.evaluate('togglePaperQuestionAnswer(9); true')
        browser.evaluate("paperAnswerTest.responses[9](new Response(JSON.stringify({answer_markdown:'$9$'}),{status:200})); true")
        browser.command('wait', '--fn', '!PaperStore.answerLoadingIds.has(9)')
        assert browser.evaluate("document.getElementById('paper-q-answer-9').hidden && !PaperStore.expandedAnswerIds.has(9)")
        assert browser.evaluate('paperAnswerTest.requests.filter(id=>id===9).length') == 2

        # A late answer on another page may populate the cache, not redraw this page.
        browser.evaluate("togglePaperQuestionAnswer(10); switchPaperStreamTab('selected').then(()=>changePaperStreamPage('selected',2)).then(()=>true)")
        browser.evaluate("paperAnswerTest.nextPage = document.querySelector('#paperQuestionStream .paper-resource-card'); paperAnswerTest.responses[10](new Response(JSON.stringify({answer_markdown:'$10$'}),{status:200})); true")
        browser.command('wait', '--fn', '!PaperStore.answerLoadingIds.has(10)')
        assert browser.evaluate("PaperStore.streamPagination.selected.page === 2 && paperAnswerTest.nextPage.isConnected && !document.getElementById('paper-q-render-10')")
    finally:
        browser.evaluate('window.fetch = paperAnswerTest.originalFetch; true')


@pytest.mark.parametrize('viewport', [(1600, 1000), (375, 812)])
def test_preview_controls_scroll_with_pages_and_metadata_has_one_toggle(browser, viewport):
    browser.command('set', 'viewport', *map(str, viewport))
    browser.evaluate((ROOT / 'tests' / 'browser_preview_fixture.js').read_text())
    browser.evaluate("selectWorkspace('paper', '智能组卷'); true")
    browser.command('wait', '--fn', '!PaperStore.cartQuestionLoad.loading && !PaperStore.streamPagination.all.loading')
    browser.evaluate('MathBankBrowserChecks.seed(12, 2); true')
    browser.settle()
    if browser.evaluate("document.getElementById('paperMetaDetails').hidden"):
        browser.command('click', '#togglePaperMetaBtn')
    browser.command('fill', '#paperMetaTitle', '预览滚动与配置保留检查')
    browser.command('click', '#togglePaperMetaBtn')
    assert browser.evaluate("document.getElementById('paperMetaDetails').hidden && document.getElementById('togglePaperMetaBtn').getAttribute('aria-expanded') === 'false'")
    browser.command('click', '#togglePaperMetaBtn')
    assert browser.evaluate("document.getElementById('paperMetaTitle').value") == '预览滚动与配置保留检查'
    assert browser.evaluate("document.querySelectorAll('#paperMetaDetails summary, #paperAiDetails summary').length") == 0
    browser.command('click', '#togglePaperMetaBtn')
    if viewport[0] < 960:
        browser.command('click', '[data-paper-pane="preview"]')
    browser.settle()
    before = browser.evaluate(r"""
    (() => {
        const sheet = document.getElementById('a4PaperPreviewSheet');
        sheet.scrollTop = 0;
        const actions = sheet.querySelector('.paper-preview-actions');
        const page = sheet.querySelector('.a4-paper-sheet');
        return {toolbarTop:actions.getBoundingClientRect().top, toolbarHeight:actions.getBoundingClientRect().height,
            pageTop:page.getBoundingClientRect().top, listScroll:document.getElementById('paperQuestionStream').scrollTop,
            rootScroll:document.scrollingElement.scrollTop};
    })()
    """)
    browser.evaluate(f"document.getElementById('a4PaperPreviewSheet').scrollTop = {before['toolbarHeight'] + 100}; true")
    browser.settle()
    after = browser.evaluate(r"""
    (() => {
        const sheet = document.getElementById('a4PaperPreviewSheet');
        const actions = sheet.querySelector('.paper-preview-actions').getBoundingClientRect();
        return {scroll:sheet.scrollTop, toolbarTop:actions.top, toolbarBottom:actions.bottom,
            viewportTop:sheet.getBoundingClientRect().top, pageTop:sheet.querySelector('.a4-paper-sheet').getBoundingClientRect().top,
            listScroll:document.getElementById('paperQuestionStream').scrollTop, rootScroll:document.scrollingElement.scrollTop,
            controlsInsidePaper:!!sheet.querySelector('.a4-paper-sheet .paper-preview-actions')};
    })()
    """)
    assert after['scroll'] > before['toolbarHeight'], (before, after)
    assert abs(before['toolbarTop'] - after['toolbarTop'] - after['scroll']) < 2, (before, after)
    assert abs(before['pageTop'] - after['pageTop'] - after['scroll']) < 2, (before, after)
    assert after['toolbarBottom'] < after['viewportTop'], after
    assert after['listScroll'] == before['listScroll'] and after['rootScroll'] == before['rootScroll'], after
    assert not after['controlsInsidePaper']
    browser.evaluate('updatePaperQuestionScore(1, 9); true')
    browser.settle()
    assert abs(browser.evaluate("document.getElementById('a4PaperPreviewSheet').scrollTop") - after['scroll']) < 2


def test_settings_offer_deepseek_flash_ocr_and_upgrade_official_aliases(browser):
    browser.command('set', 'viewport', '1440', '1000')
    browser.command('reload')
    browser.command('wait', '--fn', "typeof onModelProviderChange === 'function'")
    browser.evaluate(r"""
    (async () => {
        localStorage.setItem('custom_models_deepseek', JSON.stringify(['deepseek-v4-flash', 'deepseek-v4-flash-vision-exp']));
        const form = new FormData();
        Object.entries({
            deepseek_key:'test-deepseek-key', prefer_engine:'siliconflow',
            prefer_solve_model:'DEEPSEEK/deepseek-v4-flash',
            prefer_parse_model:'DEEPSEEK/deepseek-v4-flash',
            prefer_classify_model:'DEEPSEEK/deepseek-v4-flash',
            prefer_draw_model:'DEEPSEEK/deepseek-v4-flash'
        }).forEach(([key,value]) => form.set(key,value));
        const response = await fetch('/api/settings/save', {method:'POST', body:form});
        if (!response.ok) throw new Error('Failed to prepare isolated settings');
        return true;
    })()
    """)
    browser.command('click', '#appNavSettings')
    browser.command('wait', '--fn', "document.getElementById('settings_solve_model_select')?.value === 'deepseek-flash' && document.getElementById('settingsMetadataJson')?.value.length > 10")
    assert browser.evaluate("document.getElementById('ocrModelProvider').value") == 'siliconflow'
    for task in ('solve', 'parse', 'classify', 'draw'):
        assert browser.evaluate(f"document.getElementById('settings_{task}_model_select').value") == 'deepseek-flash'
        options = browser.evaluate(f"Array.from(document.getElementById('settings_{task}_model_select').options).map(o=>o.value)")
        assert options == ['deepseek-flash', 'deepseek-v4-pro']
    browser.command('select', '#ocrModelProvider', 'deepseek')
    assert browser.evaluate("document.getElementById('settings_ocr_model_select').value") == 'deepseek-flash'
    assert browser.evaluate("Array.from(document.getElementById('settings_ocr_model_select').options).map(o=>o.value)") == ['deepseek-flash']
    browser.command('click', '#btnSettingsSave')
    browser.command('wait', '--fn', "document.getElementById('settingsModal').classList.contains('hidden')")
    settings = browser.evaluate("fetch('/api/settings').then(r=>r.json()).then(s=>({engine:s.prefer_engine,ocr:s.deepseek_model,solve:s.prefer_solve_model}))")
    assert settings == {'engine':'deepseek', 'ocr':'deepseek-flash', 'solve':'DEEPSEEK/deepseek-flash'}
    browser.command('click', '#appNavSettings')
    browser.command('wait', '--fn', "document.getElementById('ocrModelProvider').value === 'deepseek' && document.getElementById('settings_ocr_model_select')?.value === 'deepseek-flash'")
    browser.command('click', '#settingsModal button[aria-label="关闭系统设置"]')


def test_question_without_answer_has_visible_independent_cart_action(browser):
    browser.command('set', 'viewport', '1600', '1000')
    question_id = browser.evaluate(r"""
    (async () => {
        const form = new FormData();
        form.set('content', '未提供答案也可以组卷：计算 $2+3$。');
        form.set('question_type', 'single_choice');
        form.set('difficulty', 'medium');
        form.set('answer_markdown', '');
        const response = await fetch('/api/questions', {method:'POST', body:form});
        const result = await response.json();
        if (!response.ok) throw Error(JSON.stringify(result));
        PaperStore.cart = [];
        PaperStore.filters.tab = 'all';
        PaperStore.meta.paper_type = 'exam';
        selectWorkspace('paper', '组卷排版工作台');
        await renderPaperWorkspace();
        return result.question.id;
    })()
    """)
    card = f'[data-paper-question-id="{question_id}"]'
    browser.command('wait', card)
    try:
        for theme in ('obsidian', 'violet', 'ocean', 'emerald', 'amber', 'crimson'):
            for dark in (False, True):
                browser.evaluate(f"changeTheme('{theme}', false); document.documentElement.classList.toggle('dark', {str(dark).lower()}); true")
                browser.settle()
                state = browser.evaluate(f"""
                (() => {{
                    const add = document.querySelector('{card} .paper-cart-action');
                    const answer = document.getElementById('paper-answer-toggle-{question_id}');
                    const style = getComputedStyle(add);
                    return {{disabled:add.disabled, label:add.textContent.trim(), background:style.backgroundColor,
                        color:style.color, opacity:style.opacity, visible:add.getBoundingClientRect().width > 0,
                        answerDisabled:answer.disabled, answerText:answer.textContent.trim()}};
                }})()
                """)
                assert state['answerDisabled'] and state['answerText'] == '暂无答案', state
                assert not state['disabled'] and state['visible'] and '加入试卷' in state['label'], state
                assert state['background'] not in ('transparent', 'rgba(0, 0, 0, 0)'), (theme, dark, state)
                assert state['background'] != state['color'] and float(state['opacity']) > 0.5, state
                browser.command('click', card + ' .paper-cart-action')
                browser.command('wait', '--fn', f'isInCart({question_id})')
                assert browser.evaluate(f"!!document.querySelector('{card} .paper-cart-action.is-selected')")
                assert browser.evaluate(f"PaperStore.cart.find(q => q.id === {question_id}).score") == 5
                assert browser.evaluate("document.getElementById('a4PaperPreviewSheet').textContent.includes('未提供答案也可以组卷')")
                browser.command('click', card + ' .paper-cart-action.is-selected')
                browser.command('wait', '--fn', f'!isInCart({question_id})')
        stored = browser.evaluate(f"fetch('/api/questions/{question_id}').then(r=>r.json())")
        assert stored['answer_markdown'] == ''
    finally:
        browser.evaluate("changeTheme('violet', false); document.documentElement.classList.remove('dark'); true")


def test_import_refresh_keeps_existing_question_clean_without_hiding_real_edits(browser):
    browser.command('reload')
    browser.command('set', 'viewport', '1600', '1000')
    browser.command('wait', '--fn', "typeof loadCategories === 'function' && window.categoryTree && Object.keys(categoryTree).length > 0")
    old = browser.evaluate(r"""
    (async () => {
        const book = Object.keys(categoryTree)[0];
        const chapter = Object.keys(categoryTree[book])[0];
        const content = '保留题：集合 $A=\\{1,2,3,4,5\\}$，写出它的所有单元素子集。';
        const form = new FormData();
        Object.entries({content, question_type:'detailed_answer', difficulty:'medium', category_compulsory:book, category_chapter:chapter}).forEach(([key,value])=>form.set(key,value));
        const response = await fetch('/api/questions',{method:'POST',body:form});
        const result = await response.json();
        if (!response.ok) throw Error(JSON.stringify(result));
        selectWorkspace('bank','题库管理');
        loadQuestions();
        return {id:result.question.id, content, book, chapter};
    })()
    """)
    browser.command('wait', '--fn', f"!!document.querySelector('#questionsList [data-id=\"{old['id']}\"]')")
    browser.command('click', f'#questionsList [data-id="{old["id"]}"]')
    browser.command('wait', '--fn', f'EditorState.questionId === {old["id"]} && !questionDetailLoading && !isEditorModified()')
    baseline = browser.evaluate('JSON.stringify(originalQuestionState)')
    browser.command('click', '#appNavImport')
    browser.evaluate(r"""
    (() => {
        const book = Object.keys(categoryTree)[0];
        const chapter = Object.keys(categoryTree[book])[0];
        replaceParsedQuestions([{
            content:'导入题：正方体的棱长为 $2$，求该正方体的表面积。',
            answer_markdown:'表面积为 $24$。', question_type:'detailed_answer', difficulty:'medium',
            category_compulsory:book, category_chapter:chapter, category_knowledge:'', image_paths:[], source:'导入刷新回归'
        }]);
        renderParsedQuestionsList(parsedQuestionsData);
        document.getElementById('importPlaceholder').classList.add('hidden');
        document.getElementById('parsedQuestionsWrapper').classList.remove('hidden');
        return true;
    })()
    """)
    browser.command('click', '#parsed-card-0 .card-save-btn')
    browser.command('wait', '--fn', 'parsedQuestionsData[0].saved === true')
    browser.evaluate('loadCategories().then(()=>true)')
    browser.command('click', '#appNavBank')
    state = browser.evaluate(r"""
    (() => ({id:EditorState.questionId, dirty:isEditorModified(),
        difficulty:document.getElementById('editDifficulty').value,
        baseline:JSON.stringify(originalQuestionState), content:document.getElementById('editContent').value}))()
    """)
    assert not state['dirty'], state
    assert state['id'] == old['id'] and state['content'] == old['content'], state
    assert state['baseline'] == baseline, state
    assert state['difficulty'] == 'medium', state

    # Refreshing import metadata must not mark subsequent real changes as saved.
    browser.command('click', '#editQuestionFromPreviewBtn')
    browser.command('click', '[data-editor-panel-target="content"]')
    changed = old['content'] + '\n用户补充的未保存内容。'
    browser.command('fill', '#editContent', changed)
    browser.evaluate('loadCategories().then(()=>true)')
    assert browser.evaluate('isEditorModified()')
    assert browser.evaluate("document.getElementById('editContent').value") == changed
    assert browser.evaluate('JSON.stringify(originalQuestionState)') == baseline
    browser.command('click', '#editorSection .question-editor-close')
    browser.command('wait', '--fn', "!!document.getElementById('unsavedChangesModalTitle')")
    browser.command('click', '#discardBtn')
    browser.command('wait', '--fn', "document.getElementById('editorSection').classList.contains('hidden')")
    stored = browser.evaluate(f"fetch('/api/questions/{old['id']}').then(r=>r.json()).then(q=>({{content:q.content,difficulty:q.difficulty}}))")
    assert stored == {'content':old['content'], 'difficulty':'medium'}
